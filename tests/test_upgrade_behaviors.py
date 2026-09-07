"""Behavior tests for the resume and workbench upgrade."""

import io

from services.ai_service import build_analysis_prompt, build_chat_prompt, build_matching_prompt
from services.matcher import MatchResult
from services.profile_service import resume_coverage, validate_profile
from services.resume_parser import _parse_ai_response, extract_text_from_docx, parse_resume
from ui.session import init_session_state, mark_profile_changed, select_first_job, selected_job_result


def test_injected_ai_parser_is_called_once_and_receives_full_text(monkeypatch):
    long_tail = "项目尾部仍然需要被发送给模型。" * 350
    monkeypatch.setattr("services.resume_parser.extract_text_from_pdf", lambda _: f"姓名：张三\n{long_tail}")
    calls = []

    def fake_ai(messages, temperature=0.7):
        calls.append((messages, temperature))
        return '{"name":"张三","education":"本科","school":"北大","major":"计算机","grad_year":"2026","skills":["Python"],"experience":"项目"}'

    parsed = parse_resume(b"pdf-bytes", "resume.pdf", ai_service=fake_ai)

    assert len(calls) == 1
    assert calls[0][1] == 0.1
    assert long_tail in calls[0][0][1]["content"]
    assert parsed["_parse_method"] == "ai"
    assert parsed["_coverage"] == 100


def test_docx_table_text_is_preserved():
    from docx import Document

    document = Document()
    document.add_paragraph("候选人")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "学校"
    table.cell(0, 1).text = "浙江大学"
    table.cell(1, 0).text = "技能"
    table.cell(1, 1).text = "Rust / PostgreSQL"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text_from_docx(buffer.getvalue())
    assert "浙江大学" in text
    assert "Rust / PostgreSQL" in text


def test_parser_handles_non_object_and_mixed_skill_values():
    assert _parse_ai_response("[1, 2, 3]")["skills"] == []
    parsed = _parse_ai_response('{"skills":["Python", 3, null, " python "]}')
    assert parsed["skills"] == ["Python"]


def test_local_fallback_handles_spaced_chinese_resume_layout():
    raw = """姓 名 测试
计算机技术硕士在读
南京邮电大学  计算机技术  |  硕士在读\t2025.09 - 2028.07
项目经历
独立完成 Agent CLI 项目。
专业技能
工程基础  Python（类型标注、日志与 pytest）/ Git / Docker，具备模块化开发能力。"""
    parsed = _parse_ai_response(raw)
    assert parsed["name"] == "测试"
    assert parsed["education"] == "硕士"
    assert parsed["school"] == "南京邮电大学"
    assert parsed["major"] == "计算机技术"
    assert parsed["grad_year"] == "2028"
    assert {"Python", "Git", "Docker", "pytest"}.issubset(set(parsed["skills"]))
    assert "Agent CLI" in parsed["experience"]


def test_ai_unavailable_uses_local_layout_and_marks_warning(monkeypatch):
    raw = "姓 名 测试\n南京邮电大学  计算机技术  |  硕士在读\t2025.09 - 2028.07\n项目经历\nAgent CLI 项目\n专业技能\n技术栈 Python / Git"
    monkeypatch.setattr("services.resume_parser.extract_text_from_docx", lambda _: raw)

    parsed = parse_resume(
        b"docx-bytes",
        "resume.docx",
        ai_service=lambda messages, temperature=0.1: "（AI 服务暂时不可用，请稍后重试。）",
    )

    assert parsed["_parse_method"] == "basic"
    assert parsed["_coverage"] == 100
    assert parsed["school"] == "南京邮电大学"
    assert any("AI 服务暂时不可用" in warning for warning in parsed["_warnings"])


def test_invalid_ai_json_falls_back_to_raw_text(monkeypatch):
    raw = """姓 名 测试
南京邮电大学  计算机技术  |  硕士在读	2025.09 - 2028.07
项目经历
Agent CLI 项目
专业技能
技术栈 Python / Git / Docker"""
    monkeypatch.setattr("services.resume_parser.extract_text_from_docx", lambda _: raw)

    parsed = parse_resume(
        b"docx-bytes",
        "resume.docx",
        ai_service=lambda messages, temperature=0.1: "这不是 JSON，也没有可用字段。",
    )

    assert parsed["_parse_method"] == "basic"
    assert parsed["_coverage"] == 100
    assert parsed["school"] == "南京邮电大学"
    assert any("格式异常" in warning for warning in parsed["_warnings"])


def test_ai_timeout_exception_falls_back_to_raw_text(monkeypatch):
    raw = """姓名：测试
学历：本科
学校：测试大学
专业：计算机
毕业年份：2026
技能：Python、Git
项目经历：完成一个 CLI 项目"""
    monkeypatch.setattr("services.resume_parser.extract_text_from_pdf", lambda _: raw)

    def timeout_ai(messages, temperature=0.1):
        raise TimeoutError("upstream timed out")

    parsed = parse_resume(b"pdf-bytes", "resume.pdf", ai_service=timeout_ai)

    assert parsed["_parse_method"] == "basic"
    assert parsed["_coverage"] == 100
    assert parsed["skills"] == ["Python", "Git"]
    assert any("已切换为基础文本提取" in warning for warning in parsed["_warnings"])


def test_ai_empty_object_falls_back_when_local_text_has_fields(monkeypatch):
    raw = "姓名：测试\n学校：测试大学\n专业：计算机\n毕业年份：2026\n技能：Python"
    monkeypatch.setattr("services.resume_parser.extract_text_from_pdf", lambda _: raw)

    parsed = parse_resume(
        b"pdf-bytes",
        "resume.pdf",
        ai_service=lambda messages, temperature=0.1: '{"name": null, "skills": []}',
    )

    assert parsed["_parse_method"] == "basic"
    assert parsed["_coverage"] == 71
    assert parsed["school"] == "测试大学"
    assert any("未识别到可用字段" in warning for warning in parsed["_warnings"])


def test_coverage_does_not_count_display_defaults():
    parsed = {"name": "", "education": "", "school": "", "major": "", "grad_year": "", "skills": [], "experience": "", "_recognized_fields": []}
    assert resume_coverage(parsed) == (0, 7, 0)


def test_profile_validation_is_shared_by_manual_and_resume_flows():
    assert validate_profile({"education": "本科", "school": "学校", "major": "专业", "grad_year": "2026", "skills": ["Python"]}) == []
    assert "请填写有效的四位毕业年份" in validate_profile({"education": "本科", "school": "学校", "major": "专业", "grad_year": "26", "skills": ["Python"]})


def test_session_job_selection_and_profile_change_invalidate_ai_state():
    result_a = MatchResult(job={"id": "a", "title": "岗位 A"}, score=80)
    result_b = MatchResult(job={"id": "b", "title": "岗位 B"}, score=70)
    state = {}
    init_session_state(state)
    state["matched_jobs"] = [result_a, result_b]
    select_first_job(state)
    assert selected_job_result(state) is result_a
    state["selected_job_id"] = "b"
    assert selected_job_result(state) is result_b
    state["chat_history"] = [{"role": "assistant", "content": "old"}]
    mark_profile_changed(state)
    assert state["chat_history"] == []
    assert state["analysis_cache"] == {}


def test_ai_prompts_keep_program_score_and_chat_context():
    match = MatchResult(
        job={"id": "a", "title": "后端工程师", "company": "测试公司", "requirements": ["Python"]},
        score=88.5,
        skill_match={"matched_skills": ["Python"], "missing_skills": ["Redis"]},
        reasons=["命中 1 条岗位技能要求"],
        recommendation_level="强匹配",
    )
    profile = {"name": "张三", "skills": ["Python"]}
    assert "88.5" in build_matching_prompt(profile, [match])
    assert "命中 1 条岗位技能要求" in build_analysis_prompt(profile, match.job, match)
    chat = build_chat_prompt(profile, {"matched_jobs": [match], "selected_match": match}, "第二个岗位呢？", [{"role": "user", "content": "第一个岗位怎么样？"}])
    assert "第二个岗位呢？" in chat
    assert "第一个岗位怎么样？" in chat

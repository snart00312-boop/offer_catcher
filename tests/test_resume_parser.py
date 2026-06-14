"""简历解析模块测试"""
import io
import pytest
from services.resume_parser import (
    extract_text_from_pdf,
    extract_text_from_docx,
    validate_and_fill,
    _parse_ai_response,
    ResumeParseError,
)


def _make_minimal_pdf() -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    # Use built-in CJK font so Chinese characters render correctly
    page.insert_text(
        (72, 72),
        "张三\n北京大学\n计算机科学与技术\n技能：Python, Java, SQL",
        fontname="china-s",
    )
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def _make_minimal_docx() -> bytes:
    from docx import Document
    doc = Document()
    doc.add_paragraph("李四")
    doc.add_paragraph("清华大学")
    doc.add_paragraph("软件工程")
    doc.add_paragraph("技能：Go, Docker, Kubernetes")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestPdfExtraction:
    def test_extract_text_from_valid_pdf(self):
        text = extract_text_from_pdf(_make_minimal_pdf())
        assert "张三" in text
        assert "北京大学" in text
        assert "Python" in text

    def test_extract_text_from_corrupted_pdf(self):
        with pytest.raises(ResumeParseError) as exc:
            extract_text_from_pdf(b"this is not a pdf")
        assert "PDF" in exc.value.message


class TestDocxExtraction:
    def test_extract_text_from_valid_docx(self):
        text = extract_text_from_docx(_make_minimal_docx())
        assert "李四" in text
        assert "清华大学" in text
        assert "Go" in text

    def test_extract_text_from_corrupted_docx(self):
        with pytest.raises(ResumeParseError) as exc:
            extract_text_from_docx(b"not a docx file")
        assert "DOCX" in exc.value.message


class TestValidateAndFill:
    def test_normalizes_education_values(self):
        parsed = {"name": "张三", "education": "本科", "school": "北京大学",
                   "major": "计算机", "grad_year": "2025",
                   "skills": ["Python", "Java", "SQL"], "experience": "在字节跳动实习3个月"}
        cleaned, missing = validate_and_fill(parsed)
        assert cleaned["education"] == "本科"
        assert cleaned["name"] == "张三"
        assert missing == []

    def test_maps_education_variants(self):
        cases = [("本科", "本科"), ("学士", "本科"), ("硕士", "硕士"),
                 ("硕士研究生", "硕士"), ("博士", "博士"), ("博士研究生", "博士"),
                 ("专科", "专科"), ("大专", "专科"), ("MBA", "硕士")]
        for raw, expected in cases:
            cleaned, _ = validate_and_fill({"education": raw, "name": None, "school": None,
                                              "major": None, "grad_year": None,
                                              "skills": [], "experience": None})
            assert cleaned["education"] == expected, f"{raw} -> {expected}, got {cleaned['education']}"

    def test_deduplicates_skills(self):
        parsed = {"name": "王五", "education": "本科", "school": "某大学",
                   "major": "数学", "grad_year": "2026",
                   "skills": ["Python", "python", "Java", "JAVA", " SQL "],
                   "experience": None}
        cleaned, _ = validate_and_fill(parsed)
        assert len(cleaned["skills"]) == 3

    def test_reports_missing_fields(self):
        parsed = {"name": None, "education": None, "school": None,
                   "major": "物理", "grad_year": None, "skills": [], "experience": None}
        _, missing = validate_and_fill(parsed)
        assert "education" in missing
        assert "school" in missing
        assert "skills" in missing
        assert "major" not in missing

    def test_handles_null_skills(self):
        parsed = {"name": "赵六", "education": "硕士", "school": "浙大",
                   "major": "AI", "grad_year": "2025", "skills": None, "experience": "实习"}
        cleaned, _ = validate_and_fill(parsed)
        assert cleaned["skills"] == []


class TestParseAiResponse:
    def test_parses_valid_json_response(self):
        raw = '```json\n{"name":"张三","education":"本科","school":"北大","major":"CS","grad_year":"2025","skills":["Python","Java"],"experience":"实习"}\n```'
        result = _parse_ai_response(raw)
        assert result["name"] == "张三"
        assert result["skills"] == ["Python", "Java"]

    def test_parses_json_without_code_fence(self):
        raw = '{"name":"李四","education":"硕士","school":"清华","major":"SE","grad_year":"2026","skills":["Go"],"experience":null}'
        result = _parse_ai_response(raw)
        assert result["name"] == "李四"

    def test_fallback_regex_on_invalid_json(self):
        raw = "姓名：王五\n学历：本科\n学校：复旦大学\n专业：数学\n毕业年份：2025\n技能：Python, Java, SQL\n经历：在某公司实习"
        result = _parse_ai_response(raw)
        assert result["name"] == "王五"
        assert result["school"] == "复旦大学"
        assert "Python" in result["skills"]

    def test_returns_empty_on_garbled_response(self):
        raw = "抱歉，我无法解析这份文件。"
        result = _parse_ai_response(raw)
        assert result["name"] is None
        assert result["education"] is None


class TestParseResumeIntegration:
    def test_parse_resume_pdf_without_ai(self):
        from services.resume_parser import parse_resume
        pdf_bytes = _make_minimal_pdf()
        result = parse_resume(pdf_bytes, "resume.pdf", ai_service=None)
        assert "_raw_text" in result
        assert "张三" in result["_raw_text"]
        assert isinstance(result.get("skills"), list)
